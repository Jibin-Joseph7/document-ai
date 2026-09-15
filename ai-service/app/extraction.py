"""
Turns an uploaded file on disk into plain text, dispatched by MIME type.
This is the first stage of the RAG pipeline: extraction -> chunking (commit
11) -> embeddings (commit 12) -> vector store (commit 13).
"""
import csv
import io
from pathlib import Path

from pypdf import PdfReader
from docx import Document as DocxDocument
from openpyxl import load_workbook

# Mirrors server/src/config/index.js's uploads.allowedMimeTypes - kept in
# sync manually since the two services don't share a config file format.
MIME_PDF = "application/pdf"
MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIME_TXT = "text/plain"
MIME_CSV = "text/csv"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_XLS = "application/vnd.ms-excel"  # legacy binary format, see note below


class ExtractionError(Exception):
    """Raised when a file can't be turned into text at all (corrupt,
    encrypted with no accessible text layer, unsupported format, etc.)."""


def extract_pdf(filepath: str) -> str:
    try:
        reader = PdfReader(filepath)
    except Exception as e:
        raise ExtractionError(f"Could not open PDF: {e}") from e

    if reader.is_encrypted:
        # Try an empty password first (common for "restricted but not
        # really password-protected" PDFs) before giving up.
        try:
            reader.decrypt("")
        except Exception:
            raise ExtractionError("PDF is password-protected; cannot extract text")

    pages_text = []
    for i, page in enumerate(reader.pages):
        try:
            pages_text.append(page.extract_text() or "")
        except Exception as e:
            # One malformed page shouldn't sink the whole document - note
            # it and keep going with the rest.
            pages_text.append(f"[Page {i + 1}: extraction failed - {e}]")

    text = "\n\n".join(t for t in pages_text if t.strip())
    if not text.strip():
        raise ExtractionError(
            "No extractable text found (this may be a scanned/image-only PDF "
            "that needs OCR - see the OCR bonus feature in the README)"
        )
    return text


def extract_docx(filepath: str) -> str:
    try:
        doc = DocxDocument(filepath)
    except Exception as e:
        raise ExtractionError(f"Could not open DOCX: {e}") from e

    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    # Tables carry real content in policy/handbook-style documents (leave
    # balances, approval matrices, etc.) - don't silently drop them.
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    text = "\n".join(parts)
    if not text.strip():
        raise ExtractionError("DOCX contains no extractable text")
    return text


def extract_txt(filepath: str) -> str:
    # Try UTF-8 first (the common case); fall back to latin-1, which can
    # decode any byte sequence, so this never raises on encoding alone.
    raw = Path(filepath).read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")

    if not text.strip():
        raise ExtractionError("Text file is empty")
    return text


def extract_csv(filepath: str) -> str:
    raw = Path(filepath).read_bytes()
    try:
        decoded = raw.decode("utf-8")
    except UnicodeDecodeError:
        decoded = raw.decode("latin-1")

    reader = csv.reader(io.StringIO(decoded))
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        raise ExtractionError("CSV file has no rows")

    # Render as "header: value" pairs when there's a header row plus data -
    # much more useful for embeddings/RAG than a flat comma-joined blob,
    # since each row reads like a small self-contained fact.
    header, *data_rows = rows
    lines = []
    if data_rows:
        for row in data_rows:
            pairs = [f"{h.strip()}: {v.strip()}" for h, v in zip(header, row) if v.strip()]
            lines.append(", ".join(pairs) if pairs else ", ".join(row))
    else:
        lines.append(", ".join(header))

    return "\n".join(lines)


def extract_xlsx(filepath: str) -> str:
    try:
        wb = load_workbook(filepath, read_only=True, data_only=True)
    except Exception as e:
        raise ExtractionError(f"Could not open XLSX: {e}") from e

    sheet_texts = []
    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        rows = [r for r in rows if any(c is not None and str(c).strip() for c in r)]
        if not rows:
            continue

        header, *data_rows = rows
        lines = [f"--- Sheet: {sheet.title} ---"]
        for row in data_rows:
            pairs = [
                f"{h}: {v}"
                for h, v in zip(header, row)
                if v is not None and str(v).strip()
            ]
            lines.append(", ".join(pairs) if pairs else ", ".join(str(c) for c in row if c is not None))
        sheet_texts.append("\n".join(lines))

    text = "\n\n".join(sheet_texts)
    if not text.strip():
        raise ExtractionError("XLSX file has no data in any sheet")
    return text


_EXTRACTORS = {
    MIME_PDF: extract_pdf,
    MIME_DOCX: extract_docx,
    MIME_TXT: extract_txt,
    MIME_CSV: extract_csv,
    MIME_XLSX: extract_xlsx,
    # MIME_XLS (legacy .xls / BIFF format) is intentionally not supported -
    # openpyxl can't read it, and adding xlrd just for the deprecated
    # binary format isn't worth the extra dependency for this project.
    # Ask users to re-save as .xlsx.
}


def extract_text(filepath: str, mime_type: str) -> str:
    """Extract plain text from `filepath` based on its `mime_type`.
    Raises ExtractionError for unsupported types or unreadable content.
    """
    extractor = _EXTRACTORS.get(mime_type)
    if extractor is None:
        raise ExtractionError(f"Unsupported mime type for extraction: {mime_type}")
    return extractor(filepath)